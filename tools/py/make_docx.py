"""Builds a .docx from a JSON spec on stdin. Usage: python make_docx.py <workspace>

{"path": "report.docx", "title": "Test report", "blocks": [
   {"heading": "Step 1", "level": 1}, {"text": "A paragraph."}, {"bullets": ["a", "b"]}, {"numbered": ["a", "b"]},
   {"image": "screenshots/step-01.png", "width": 6, "caption": "Clicked Login"},
   {"table": [["Step", "Result"], ["1", "Pass"]]}, {"page_break": true}]}
Paths are relative to the workspace and must stay inside it.
"""
import json
import os
import sys

from docx import Document
from docx.shared import Inches


def inside(workspace, p):
    full = os.path.abspath(os.path.join(workspace, p))
    if os.path.commonpath([full, workspace]) != workspace:
        raise ValueError(f"{full} is outside the workspace")
    return full


def main():
    workspace = os.path.abspath(sys.argv[1])
    spec = json.loads(sys.stdin.read())
    out = inside(workspace, spec.get("path") or "document.docx")
    if not out.lower().endswith(".docx"):
        out += ".docx"
    doc = Document()
    if spec.get("title"):
        doc.add_heading(str(spec["title"]), level=0)
    for b in spec.get("blocks", []):
        if "heading" in b:
            doc.add_heading(str(b["heading"]), level=max(1, min(int(b.get("level", 1)), 4)))
        elif "text" in b:
            doc.add_paragraph(str(b["text"]))
        elif "bullets" in b:
            for item in b["bullets"]:
                doc.add_paragraph(str(item), style="List Bullet")
        elif "numbered" in b:
            for item in b["numbered"]:
                doc.add_paragraph(str(item), style="List Number")
        elif "image" in b:
            doc.add_picture(inside(workspace, b["image"]), width=Inches(float(b.get("width", 6))))
            if b.get("caption"):
                doc.add_paragraph(str(b["caption"]), style="Caption")
        elif "table" in b:
            rows = b["table"]
            if rows:
                t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                t.style = "Table Grid"
                for i, r in enumerate(rows):
                    for j, cell in enumerate(r):
                        t.cell(i, j).text = str(cell)
        elif b.get("page_break"):
            doc.add_page_break()
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    print(f"wrote {os.path.relpath(out, workspace)} ({len(spec.get('blocks', []))} blocks)")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:  # report to the agent instead of a traceback
        print(f"error: {e}")
        sys.exit(1)
