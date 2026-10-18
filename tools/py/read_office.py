"""Prints the text of a .docx, .pptx or .xlsx file, so the agent can check a document it made.
Usage: python read_office.py <file>

docx: paragraphs in order ([Heading 1] / [Caption] ... style names), [image] where a picture sits, tables as rows.
pptx: one block per slide with its texts, [picture] markers and speaker notes.
xlsx: one block per sheet, rows as "a | b | c" (first 200 rows).
"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")

PIC = "{http://schemas.openxmlformats.org/drawingml/2006/picture}pic"


def docx(path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = Document(path)
    out = []
    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, d)
            style = p.style.name if p.style is not None else ""
            mark = f"[{style}] " if style and style != "Normal" else ""
            pics = len(child.findall(".//" + PIC))
            text = p.text.strip()
            if pics:
                out.append(f"{mark}[image]" * pics + (f" {text}" if text else ""))
            elif text:
                out.append(mark + text)
        elif tag == "tbl":
            out.append("[table]")
            for row in Table(child, d).rows:
                out.append(" | ".join(c.text.strip() for c in row.cells))
    out.append(f"-- {len(d.inline_shapes)} picture(s), {sum(1 for p in d.paragraphs if p.style is not None and p.style.name.startswith('Heading'))} heading(s)")
    return out


def pptx(path):
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        out.append(f"--- slide {i}")
        for sh in slide.shapes:
            if sh.shape_type == 13:  # picture
                out.append("[picture]")
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        out.append(t)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("[notes] " + slide.notes_slide.notes_text_frame.text.strip())
    return out


def xlsx(path):
    import openpyxl

    out = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        out.append(f"--- sheet {ws.title}")
        for row in ws.iter_rows(values_only=True, max_row=200):
            out.append(" | ".join("" if v is None else str(v) for v in row))
    return out


def main():
    path = sys.argv[1]
    ext = os.path.splitext(path)[1].lower()
    read = {".docx": docx, ".pptx": pptx, ".xlsx": xlsx}.get(ext)
    if not read:
        sys.exit(f"not a .docx, .pptx or .xlsx file: {path}")
    print("\n".join(read(path)))


if __name__ == "__main__":
    main()
