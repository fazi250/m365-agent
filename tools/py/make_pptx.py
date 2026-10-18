"""Builds a .pptx from a JSON spec on stdin. Usage: python make_pptx.py <workspace>

{"path": "deck.pptx", "title": "Test run", "subtitle": "17 Oct 2026", "slides": [
   {"title": "Step 1", "bullets": ["Opened the page", "Clicked Login"], "image": "screenshots/step-01.png", "notes": "..."}]}
A slide with bullets and an image puts the text left and the picture right. Paths are relative to the workspace.
"""
import json
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches, Pt


def inside(workspace, p):
    full = os.path.abspath(os.path.join(workspace, p))
    if os.path.commonpath([full, workspace]) != workspace:
        raise ValueError(f"{full} is outside the workspace")
    return full


def add_picture_fit(slide, path, left, top, max_w, max_h):
    with Image.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    pw, ph = int(w * scale), int(h * scale)
    slide.shapes.add_picture(path, left + (max_w - pw) // 2, top + (max_h - ph) // 2, Emu(pw), Emu(ph))


def main():
    workspace = os.path.abspath(sys.argv[1])
    spec = json.loads(sys.stdin.read())
    out = inside(workspace, spec.get("path") or "presentation.pptx")
    if not out.lower().endswith(".pptx"):
        out += ".pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    W, H = prs.slide_width, prs.slide_height
    if spec.get("title"):
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = str(spec["title"])
        s.placeholders[1].text = str(spec.get("subtitle", ""))
    for d in spec.get("slides", []):
        s = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        top, margin = Inches(1.5), Inches(0.5)
        title = s.shapes.title
        title.text = str(d.get("title", ""))
        title.left, title.top, title.width, title.height = margin, Inches(0.3), W - 2 * margin, Inches(1.0)
        bullets, image = d.get("bullets") or [], d.get("image")
        text_w = (W - 3 * margin) // 2 if image and bullets else W - 2 * margin
        if bullets:
            tf = s.shapes.add_textbox(margin, top, text_w, H - top - margin).text_frame
            tf.word_wrap = True
            for i, line in enumerate(bullets):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.text = f"• {line}"
                para.font.size = Pt(20)
        if image:
            left = margin + text_w + margin if bullets else margin
            add_picture_fit(s, inside(workspace, image), left, top, W - left - margin, H - top - margin)
        if d.get("notes"):
            s.notes_slide.notes_text_frame.text = str(d["notes"])
    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)
    print(f"wrote {os.path.relpath(out, workspace)} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"error: {e}")
        sys.exit(1)
